"""
Excel compliance matrix and financial schedule generator using openpyxl.
"""
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelGenerator:
    """Creates formatted Excel spreadsheets for technical compliance and turnover schedules."""

    @classmethod
    def generate_compliance_sheet(
        cls,
        output_path: Path,
        tender_number: str,
        buyer_name: str,
        company_name: str,
        parameters: List[Dict[str, str]]
    ) -> Path:
        """
        Creates a technical compliance matrix workbook.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Technical Compliance"

        # Theme styles
        navy_header_fill = PatternFill(start_color="102C57", end_color="102C57", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="102C57")
        sub_font = Font(name="Calibri", size=10, italic=True)
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)

        thin_border = Border(
            left=Side(style='thin', color="D3D3D3"),
            right=Side(style='thin', color="D3D3D3"),
            top=Side(style='thin', color="D3D3D3"),
            bottom=Side(style='thin', color="D3D3D3")
        )

        # Title Block
        ws.merge_cells("A1:E1")
        ws["A1"] = f"TECHNICAL COMPLIANCE STATEMENT"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:E2")
        ws["A2"] = f"Tender / Bid No: {tender_number} | Buyer: {buyer_name} | Bidder: {company_name}"
        ws["A2"].font = sub_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        # Table Header
        headers = ["S.No.", "Clause / Parameter", "Buyer Specification Requirement", "Offered Specification / Make & Model", "Compliance (Yes/No)", "Deviation / Remarks"]
        ws.row_dimensions[4].height = 28

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = navy_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Populate rows
        if not parameters:
            parameters = [
                {"param": "General Technical Specifications", "buyer_spec": "As per Tender Schedule of Requirements", "offered": "Fully Complied as per catalog", "compliance": "Yes", "remarks": "No deviations"},
                {"param": "Warranty Period", "buyer_spec": "Comprehensive Onsite Warranty as specified in ATC", "offered": "Comprehensive Onsite OEM Warranty", "compliance": "Yes", "remarks": "Full warranty support"},
                {"param": "Standards & Certifications", "buyer_spec": "ISO / BIS / CE certification as specified", "offered": "Standard Certified Product Offered", "compliance": "Yes", "remarks": "Valid certificates attached"}
            ]

        compliance_green_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")

        for r_idx, item in enumerate(parameters, start=5):
            ws.cell(row=r_idx, column=1, value=r_idx - 4).alignment = Alignment(horizontal="center")
            ws.cell(row=r_idx, column=2, value=item.get("param", ""))
            ws.cell(row=r_idx, column=3, value=item.get("buyer_spec", ""))
            ws.cell(row=r_idx, column=4, value=item.get("offered", ""))

            comp_cell = ws.cell(row=r_idx, column=5, value=item.get("compliance", "Yes"))
            comp_cell.alignment = Alignment(horizontal="center")
            comp_cell.fill = compliance_green_fill
            comp_cell.font = bold_font

            ws.cell(row=r_idx, column=6, value=item.get("remarks", "Complied without deviation"))

            for c in range(1, 7):
                ws.cell(row=r_idx, column=c).border = thin_border
                ws.cell(row=r_idx, column=c).font = regular_font

        # Adjust column widths
        col_widths = {1: 8, 2: 24, 3: 40, 4: 40, 5: 16, 6: 28}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        return output_path
