"""
Tests for DOCX template preservation and Excel compliance matrix generation.
"""
from pathlib import Path
import docx
import openpyxl
from generation.docx_generator import DocxGenerator
from generation.excel_generator import ExcelGenerator

def test_template_preservation_and_replacement(tmp_path):
    template_file = tmp_path / "company_undertaking_template.docx"
    doc = docx.Document()
    doc.add_heading("OFFICIAL COMPANY UNDERTAKING", level=1)
    doc.add_paragraph("This undertaking is submitted by {{company_name}} for Tender No: {{tender_number}}.")
    doc.add_paragraph("Authorized Signatory: {{signatory_name}}")
    doc.save(str(template_file))

    out_file = tmp_path / "generated_undertaking.docx"
    replacements = {
        "{{company_name}}": "Acme Systems India Ltd",
        "{{tender_number}}": "GEM/2024/B/55555",
        "{{signatory_name}}": "Rajesh Sharma"
    }

    DocxGenerator.generate_from_template(
        template_path=template_file,
        output_path=out_file,
        replacements=replacements
    )

    assert out_file.exists()
    # Check generated content
    gen_doc = docx.Document(str(out_file))
    full_text = " ".join(p.text for p in gen_doc.paragraphs)
    assert "Acme Systems India Ltd" in full_text
    assert "GEM/2024/B/55555" in full_text
    assert "Rajesh Sharma" in full_text
    # Template itself should remain unmodified
    orig_doc = docx.Document(str(template_file))
    assert "{{company_name}}" in " ".join(p.text for p in orig_doc.paragraphs)

def test_standard_document_generation(tmp_path):
    out_file = tmp_path / "standard_doc.docx"
    DocxGenerator.generate_standard_document(
        output_path=out_file,
        title="Non-Blacklisting Declaration",
        tender_number="GEM/2024/B/112233",
        buyer_name="Ministry of Defence",
        company_name="Acme Systems India Ltd",
        body_sections=[
            {"heading": "Declaration", "paragraphs": ["We hereby declare that we are not blacklisted."]}
        ],
        signature_block="For Acme Systems\nAuthorized Signatory"
    )
    assert out_file.exists()
    doc = docx.Document(str(out_file))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "NON-BLACKLISTING DECLARATION" in text
    assert "GEM/2024/B/112233" in text

def test_excel_compliance_generation(tmp_path):
    out_file = tmp_path / "compliance_matrix.xlsx"
    ExcelGenerator.generate_compliance_sheet(
        output_path=out_file,
        tender_number="GEM/2024/B/112233",
        buyer_name="Ministry of Railways",
        company_name="Acme Systems India Ltd",
        parameters=[
            {"param": "RAM Capacity", "buyer_spec": "128 GB ECC", "offered": "128 GB ECC DDR5", "compliance": "Yes", "remarks": "Fully Complied"}
        ]
    )
    assert out_file.exists()
    wb = openpyxl.load_workbook(str(out_file))
    ws = wb.active
    assert "TECHNICAL COMPLIANCE" in ws["A1"].value
    assert ws["C5"].value == "128 GB ECC"
    assert ws["D5"].value == "128 GB ECC DDR5"
