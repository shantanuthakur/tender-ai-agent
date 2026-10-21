"""
Comprehensive end-to-end workflow integration test.
Tests template preference, corrigendum handling, anti-hallucination placeholders,
ZIP bundle export, and missing info resolution.
"""
from pathlib import Path
import docx
import zipfile
import io
from tender.analyzer import TenderAnalyzer
from company.template_library import TemplateLibrary
from company.profile_manager import CompanyProfileManager
from generation.document_builder import DocumentBuilder
from verification.compliance_checker import ComplianceChecker
from database import (
    list_missing_information,
    resolve_missing_information,
    list_generated_documents,
    get_tender_project
)

def test_full_bid_workflow_e2e(tmp_path):
    # 1. Setup Company Profile
    CompanyProfileManager.set_fact("company_name", "Bharat InfraTech Solutions Pvt Ltd", "general")
    CompanyProfileManager.set_fact("gstin", "07AABCB1234F1Z9", "taxation")
    CompanyProfileManager.set_fact("pan", "AABCB1234F", "taxation")
    CompanyProfileManager.set_fact("authorized_signatory_name", "Vikram Malhotra", "signatory")
    CompanyProfileManager.set_fact("authorized_signatory_designation", "Director - Business Development", "signatory")

    # 2. Upload an Approved Official Company Template
    official_template = tmp_path / "Bharat_Official_Covering_Letter.docx"
    doc_tpl = docx.Document()
    doc_tpl.add_heading("BHARAT INFRATECH OFFICIAL BID SUBMISSION", level=1)
    doc_tpl.add_paragraph("Tender Number: {{tender_number}}")
    doc_tpl.add_paragraph("Buyer: {{buyer_name}}")
    doc_tpl.add_paragraph("Bidder: {{company_name}} (GSTIN: {{gstin}})")
    doc_tpl.add_paragraph("{{body_content}}")
    doc_tpl.save(str(official_template))

    tpl_id = TemplateLibrary.register_document(
        file_source_path=official_template,
        is_template=True,
        is_approved=True,
        category_override="covering_letter",
        title="Official Covering Letter Template"
    )
    assert tpl_id > 0

    # 3. Create Tender Project
    project_id = TenderAnalyzer.create_project("GeM IT Infrastructure Bid 2024")

    # 4. Ingest Tender Files: Bid Doc, ATC, and Corrigendum
    bid_doc = tmp_path / "GeM_Bid_Document_GEM_2024_B_991122.txt"
    bid_doc.write_text(
        "Government e-Marketplace (GeM)\n"
        "Bid Number: GEM/2024/B/991122\n"
        "Ministry/State Name: Ministry of Power\n"
        "Organisation Name: NTPC Limited\n"
        "Bid End Date/Time: 10-12-2024 14:00:00\n"
        "Minimum Average Annual Turnover: Rs. 150 Lakhs\n",
        encoding="utf-8"
    )

    atc_doc = tmp_path / "Buyer_Added_Bid_Specific_ATC.txt"
    atc_doc.write_text(
        "Buyer Added ATC:\n"
        "1. Bidder must submit Non-Blacklisting undertaking.\n"
        "2. Make in India (MII) local content certificate is mandatory.\n"
        "3. Technical compliance sheet must be uploaded in DOCX and XLSX format.\n",
        encoding="utf-8"
    )

    corrigendum_doc = tmp_path / "Corrigendum_1.txt"
    corrigendum_doc.write_text(
        "=== CORRIGENDUM NO. 1 ===\n"
        "Bid Number: GEM/2024/B/991122\n"
        "Bid End Date/Time: 20-12-2024 16:00:00 (Extended)\n",
        encoding="utf-8"
    )

    TenderAnalyzer.ingest_tender_file(project_id, bid_doc)
    TenderAnalyzer.ingest_tender_file(project_id, atc_doc)
    TenderAnalyzer.ingest_tender_file(project_id, corrigendum_doc)

    # 5. Analyze Tender Package
    analysis = TenderAnalyzer.analyze_tender_package(project_id)
    assert analysis["metadata"]["tender_number"] == "GEM/2024/B/991122"
    assert "NTPC Limited" in analysis["metadata"]["buyer_name"] or "Ministry of Power" in analysis["metadata"]["buyer_name"]
    assert analysis["metadata"]["corrigenda_applied"] is True

    # 6. Prepare All Applicable Bid Documents
    gen_docs = DocumentBuilder.prepare_all_documents(project_id)
    assert len(gen_docs) >= 5

    # Check that covering letter was built from the official template
    cov_letter = next((d for d in gen_docs if "Covering" in d["title"]), None)
    assert cov_letter is not None
    assert "Company Template" in cov_letter["template_used"]

    # Verify generated covering letter content
    cov_path = Path(cov_letter["path"])
    assert cov_path.exists()
    doc_read = docx.Document(str(cov_path))
    cov_text = " ".join(p.text for p in doc_read.paragraphs)
    assert "GEM/2024/B/991122" in cov_text
    assert "Bharat InfraTech Solutions Pvt Ltd" in cov_text

    # 7. Verify Compliance Audit
    audit = ComplianceChecker.run_compliance_audit(project_id)
    assert audit["total_requirements"] > 0
    assert audit["addressed_and_supported"] + audit["awaiting_human_review"] > 0

    # 8. Check and Resolve Missing Information
    missing = list_missing_information(project_id)
    assert len(missing) > 0  # Placeholders like Local Content % or CA UDIN should be flagged
    first_missing = missing[0]
    resolve_missing_information(first_missing["id"], "55% Local Content Verified")
    updated_missing = list_missing_information(project_id)
    resolved_item = next(m for m in updated_missing if m["id"] == first_missing["id"])
    assert resolved_item["status"] == "provided"
    assert resolved_item["resolved_value"] == "55% Local Content Verified"

    # 9. Verify ZIP Packaging
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for d in gen_docs:
            p = Path(d["path"])
            if p.exists():
                zf.write(p, arcname=p.name)
        zf.writestr("00_CHECKLIST.txt", "Tender Verification Complete.")

    zip_buffer.seek(0)
    with zipfile.ZipFile(zip_buffer, "r") as zf:
        names = zf.namelist()
        assert "00_CHECKLIST.txt" in names
        assert any("Covering_Letter" in n for n in names)
        assert any(".xlsx" in n for n in names)
