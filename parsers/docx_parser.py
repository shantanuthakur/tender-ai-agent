"""
Word DOCX parser and structural analyzer for templates and historical tender submissions.
"""
from pathlib import Path
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor

class DocxParser:
    """Extracts text and structural styling from Word documents."""

    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        result = {
            "text": "",
            "paragraphs": [],
            "headings": [],
            "tables": [],
            "headers": [],
            "footers": [],
            "signature_blocks": [],
            "detected_styles": {
                "fonts": set(),
                "margins": {},
                "has_tables": False,
                "has_headers": False,
                "has_footers": False,
                "has_letterhead": False
            },
            "error": None
        }

        try:
            doc = docx.Document(str(file_path))

            # Analyze sections (headers, footers, margins)
            for s_idx, section in enumerate(doc.sections):
                margin_info = {
                    "top_in": round(section.top_margin.inches, 2) if section.top_margin else None,
                    "bottom_in": round(section.bottom_margin.inches, 2) if section.bottom_margin else None,
                    "left_in": round(section.left_margin.inches, 2) if section.left_margin else None,
                    "right_in": round(section.right_margin.inches, 2) if section.right_margin else None,
                }
                result["detected_styles"]["margins"] = margin_info

                # Header
                if section.header and section.header.paragraphs:
                    header_text = "\n".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
                    if header_text:
                        result["headers"].append(header_text)
                        result["detected_styles"]["has_headers"] = True
                        if any(term in header_text.lower() for term in ["ltd", "pvt", "limited", "company", "iso", "gstin", "cin", "office", "road", "tel"]):
                            result["detected_styles"]["has_letterhead"] = True

                # Footer
                if section.footer and section.footer.paragraphs:
                    footer_text = "\n".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
                    if footer_text:
                        result["footers"].append(footer_text)
                        result["detected_styles"]["has_footers"] = True

            # Parse paragraphs
            full_text_list: List[str] = []
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if not p_text:
                    continue

                full_text_list.append(p_text)

                # Collect fonts from runs
                for run in p.runs:
                    if run.font.name:
                        result["detected_styles"]["fonts"].add(run.font.name)

                # Headings
                if p.style and ("Heading" in p.style.name or "Title" in p.style.name):
                    result["headings"].append({
                        "level": p.style.name,
                        "text": p_text
                    })

                # Check for signature block indicators
                lower_p = p_text.lower()
                sig_keywords = [
                    "authorized signatory", "authorised signatory", "signature of bidder",
                    "for and on behalf of", "name and designation", "date and seal",
                    "stamp & signature", "name:", "designation:"
                ]
                if any(kw in lower_p for kw in sig_keywords):
                    result["signature_blocks"].append(p_text)

            # Parse tables
            for t_idx, table in enumerate(doc.tables):
                result["detected_styles"]["has_tables"] = True
                table_data = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_cells)

                result["tables"].append({
                    "table_index": t_idx + 1,
                    "rows_count": len(table_data),
                    "cols_count": len(table_data[0]) if table_data else 0,
                    "sample_rows": table_data[:5]
                })

                # Append table text to overall text
                table_txt = "\n".join([" | ".join(row) for row in table_data])
                full_text_list.append(f"[Table {t_idx + 1}]\n{table_txt}")

            result["text"] = "\n\n".join(full_text_list)
            result["detected_styles"]["fonts"] = sorted(list(result["detected_styles"]["fonts"]))

            # Letterhead fallback check in top paragraphs
            if not result["detected_styles"]["has_letterhead"] and len(full_text_list) > 0:
                top_block = " ".join(full_text_list[:3]).lower()
                if any(term in top_block for term in ["gstin", "cin", "regd. office", "registered office", "iso 9001"]):
                    result["detected_styles"]["has_letterhead"] = True

        except Exception as ex:
            result["error"] = f"Failed to extract DOCX: {str(ex)}"

        return result
